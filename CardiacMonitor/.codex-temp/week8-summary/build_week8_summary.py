from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE = Path(r"D:\metrial internship\Week_4_Authentication_Security_Summary.docx")
OUTPUT = Path(r"D:\Visual Studio\Mohammad-Abdelfattah-BinX-Backend-Internship\CardiacMonitor\Week_8_Advanced_Queries_Performance_Summary.docx")

BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x5B, 0x65, 0x73)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_title(doc, text):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(4)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge_name in ("top", "left", "bottom", "right", "between", "bar"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")
        p_bdr.append(edge)
    r = p.add_run(text)
    set_font(r, size=23, bold=True, color=BLACK)
    return p


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(text)
    set_font(r, size=13, color=GRAY)
    return p


def add_metadata(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    lead = p.add_run(f"{label}: ")
    set_font(lead, size=10.5, bold=True, color=BLUE)
    value_run = p.add_run(value)
    set_font(value_run, size=10.5, color=BLACK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.widow_control = True
    return p


def add_lead_paragraph(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    r1 = p.add_run(lead)
    set_font(r1, bold=True, color=BLUE)
    r2 = p.add_run(text)
    set_font(r2, color=BLACK)
    return p


def add_list_item(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1666667
    p.paragraph_format.widow_control = True
    p.add_run(text)
    return p


def replace_cell_text(cell, text):
    p = cell.paragraphs[0]
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def append_daily_table(doc, source_table):
    table_xml = deepcopy(source_table._tbl)
    body = doc._body._element
    body.insert(len(body) - 1, table_xml)
    table = doc.tables[-1]
    rows = [
        ("Day", "Focus", "Key Outcome"),
        ("1", "N+1 Diagnosis", "Sprint planning, EF Core query logging, realistic seed data, and a measured baseline."),
        ("2", "Query Optimization", "Eager loading, projection, split queries, and verified query-count reduction."),
        ("3", "Redis Caching", "Cache-aside integration, expiration policy, write-side invalidation, and hit/miss timing."),
        ("4", "Indexing & Profiling", "Targeted composite indexes, execution-plan analysis, and mentor code review."),
        ("5", "Benchmark Demo", "Sprint review, before/after evidence, iteration backlog, and retrospective."),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            replace_cell_text(cell, value)
    return table


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build():
    doc = Document(REFERENCE)
    source_table = deepcopy(doc.tables[0]._tbl)
    clear_body(doc)

    doc.core_properties.title = "Week 8 Advanced Queries and Performance Summary"
    doc.core_properties.subject = "Backend Development Internship Week 8 Summary"
    doc.core_properties.author = "BinX Tech"

    add_title(doc, "WEEK 8 SUMMARY")
    add_subtitle(doc, "Advanced Queries, Redis Caching, and Performance Tuning in ASP.NET Core")
    add_metadata(doc, "Program", "Backend Development Internship (.NET)")
    add_metadata(doc, "Week", "Week 8 of 10 | Phase 3: Applied Project Work, Sprint 3")
    add_metadata(doc, "Theme", "Advanced EF Core Queries, Redis Caching, and Measured Performance Tuning")
    add_metadata(doc, "Duration", "40 hours across five training days")

    add_heading(doc, "Executive Summary")
    add_body(doc, "Week 8 focused on making the capstone API measurably faster through better Entity Framework Core queries, Redis caching, and targeted database indexing. The central outcome was a performance-oriented API that avoids unnecessary database round-trips, serves suitable high-read data from a distributed cache, and uses indexes that match real filtering and sorting patterns.")
    add_body(doc, "The sprint treated performance as an evidence-based engineering task rather than an assumption. Interns established realistic baselines, diagnosed the N+1 query problem through SQL logging, selected eager loading, projection, or split queries according to each endpoint, tested cache invalidation on every related write path, and compared query counts, response times, and execution plans before and after each optimization.")

    add_heading(doc, "Learning Objectives Achieved")
    objectives = [
        "Diagnosed N+1 query behavior using EF Core SQL logging and realistic seeded data.",
        "Reduced unnecessary database round-trips through eager loading, projection, and split queries.",
        "Integrated Redis through IDistributedCache and applied the cache-aside pattern to suitable data.",
        "Implemented write-side cache invalidation so cached results remain consistent after changes.",
        "Added targeted and composite database indexes based on actual query patterns.",
        "Demonstrated performance improvements with before-and-after query counts, timings, and execution plans.",
    ]
    for item in objectives:
        add_list_item(doc, item)

    add_heading(doc, "Five-Day Curriculum and Practical Work")
    source_doc = Document(REFERENCE)
    append_daily_table(doc, source_doc.tables[0])

    add_heading(doc, "Day 1: Sprint Planning and N+1 Diagnosis", 2)
    add_body(doc, "Sprint 3 began with a measurable goal tied to the capstone's real endpoints. Instead of defining success as simply making the API faster, the backlog used concrete targets such as reducing the query count of an order-history endpoint from more than twenty database calls to one or two. The plan also carried forward one improvement from the Sprint 2 retrospective.")
    add_body(doc, "The N+1 problem was examined as one query for a parent list followed by one additional query for each returned item. EF Core query logging made the generated SQL visible, while realistic seed data exposed behavior that would remain hidden with only a few test records. List endpoints that accessed navigation properties inside loops were treated as the strongest candidates for investigation.")
    add_lead_paragraph(doc, "Hands-on result: ", "Enabled EF Core query logging, seeded at least 50 realistic records, measured the query counts of key list endpoints, and documented a genuine N+1 issue as a Sprint 3 backlog task.")

    add_heading(doc, "Day 2: Query Optimization with Eager Loading and Projection", 2)
    add_body(doc, "Eager loading with Include and ThenInclude was used when an endpoint genuinely needed related entities. For list and summary responses that required only selected fields, projection with Select produced leaner SQL and avoided loading complete entity graphs. When multiple collection navigations risked a cartesian-product explosion, AsSplitQuery provided a controlled tradeoff between result size and additional round-trips.")
    add_body(doc, "Each change was verified by rerunning the same endpoint under the same data conditions. A successful fix showed the query count collapse from N+1 to a small, stable number and preserved the endpoint's functional response. The choice among Include, projection, and split queries was based on the shape of the required data rather than a single rule applied everywhere.")
    for item in [
        "Use Include for detail views that require complete related entities.",
        "Prefer projection for list or summary endpoints that need only a limited set of fields.",
        "Use split queries when multiple included collections would create excessive duplicated rows.",
    ]:
        add_list_item(doc, item)

    add_heading(doc, "Day 3: Redis Caching and Cache Invalidation", 2)
    add_body(doc, "Redis was introduced for data that is read frequently and changes infrequently, such as a product catalog or category tree. ASP.NET Core's IDistributedCache abstraction kept cache access behind a framework interface, while the cache-aside pattern checked Redis first, loaded from the database on a miss, stored the serialized result with an expiration, and then returned it to the caller.")
    add_body(doc, "Cache correctness received the same attention as cache speed. Every create, update, or delete operation that affected a cached result removed or refreshed the relevant cache entry. This prevented stale responses from surviving until expiration and avoided using caching on rapidly changing data where invalidation cost would outweigh the benefit.")
    add_lead_paragraph(doc, "Hands-on result: ", "Configured Redis, cached one high-read endpoint, verified immediate freshness after related writes, and recorded the response-time difference between a cache miss and a cache hit.")

    add_heading(doc, "Day 4: Database Indexing and Performance Profiling", 2)
    add_body(doc, "Indexing decisions targeted columns used repeatedly in WHERE clauses, JOIN conditions, and ORDER BY operations. A composite index was added where a common query filtered and sorted on multiple columns together, such as CustomerId and OrderDate. Indexes were not added indiscriminately because each additional index increases storage and write overhead.")
    add_body(doc, "EF Core migrations captured the schema changes, and database execution plans showed whether the optimizer moved from an expensive scan to an index seek. The mid-sprint pull request included the reasoning for each query, caching, and indexing choice together with before-and-after evidence for mentor review.")
    add_lead_paragraph(doc, "Review standard: ", "An optimization counted as complete only when measurements or execution plans showed a real improvement and existing endpoint behavior remained correct.")

    add_heading(doc, "Day 5: Sprint Review, Benchmark Demo, and Retrospective", 2)
    add_body(doc, "The Sprint Review presented concrete results: query counts before and after eager loading or projection, cache miss and hit timings, and execution plans before and after indexing. Each backlog item was checked against its stated target, the cache invalidation path was demonstrated, and completed work was linked to a reviewed and merged pull request.")
    add_body(doc, "Unfinished performance opportunities moved into the Sprint 4 backlog rather than being treated as complete. The retrospective recorded what worked, what should improve, and one specific action for the next sprint, such as adding an automated regression test that detects a reintroduced N+1 query.")
    add_lead_paragraph(doc, "Hands-on result: ", "Delivered the benchmark demo, documented remaining opportunities, completed the Sprint 3 retrospective, and assembled the mentor check-in summary with measurements, caching strategy, and pull-request evidence.")

    add_heading(doc, "Required Deliverables")
    deliverables = [
        "A documented N+1 diagnosis with before-and-after query counts and an appropriate EF Core fix.",
        "Redis cache-aside behavior on at least one high-read endpoint with tested write-side invalidation.",
        "At least one targeted composite index supported by before-and-after profiling evidence.",
        "A reviewed and merged pull request containing the performance changes and measurements.",
        "A Sprint 3 retrospective with one concrete action for Sprint 4 testing, documentation, or deployment.",
    ]
    for item in deliverables:
        add_list_item(doc, item)

    add_heading(doc, "Performance Verification Checklist")
    checklist = [
        "Run the same endpoints against realistic seeded data and record baseline query counts and response times.",
        "Confirm that optimized endpoints return the same functional results while using the intended query strategy.",
        "Test Redis cache misses, cache hits, expiration, and every related create, update, and delete invalidation path.",
        "Compare execution plans or timings before and after each index and verify that write overhead remains justified.",
        "Attach the evidence to the pull request, demo it during the Sprint Review, and convert suitable checks into regression tests.",
    ]
    for item in checklist:
        add_list_item(doc, item, numbered=True)

    set_update_fields(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
